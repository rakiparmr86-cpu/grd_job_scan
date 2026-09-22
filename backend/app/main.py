from __future__ import annotations

import io
import uuid
from pathlib import Path
from typing import List

import cv2
import numpy as np
import pytesseract
from docx import Document
from docx.shared import Pt
from fastapi import Depends, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel

from app.auth import current_user, router as auth_router

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
TEXT_EXTENSIONS = {".txt", ".csv", ".md", ".json", ".log"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"}

BASE_DIR = Path(__file__).resolve().parent.parent
SCAN_DIR = BASE_DIR / "data" / "scans"
EXPORT_DIR = BASE_DIR / "data" / "exports"
SCAN_DIR.mkdir(parents=True, exist_ok=True)
EXPORT_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(
    title="GRD Job Scan API",
    version="1.0.0",
    description="Document scanning, enhancement, OCR and Word export API.",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Restrict this in production.
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.include_router(auth_router)


class ExportPage(BaseModel):
    scan_id: str
    text: str


class ExportRequest(BaseModel):
    title: str = "GRD Job Scan"
    pages: List[ExportPage]


def order_points(pts: np.ndarray) -> np.ndarray:
    rect = np.zeros((4, 2), dtype="float32")
    s = pts.sum(axis=1)
    diff = np.diff(pts, axis=1).reshape(-1)

    rect[0] = pts[np.argmin(s)]      # top-left
    rect[2] = pts[np.argmax(s)]      # bottom-right
    rect[1] = pts[np.argmin(diff)]   # top-right
    rect[3] = pts[np.argmax(diff)]   # bottom-left
    return rect


def four_point_transform(image: np.ndarray, pts: np.ndarray) -> np.ndarray:
    rect = order_points(pts)
    tl, tr, br, bl = rect

    width_a = np.linalg.norm(br - bl)
    width_b = np.linalg.norm(tr - tl)
    max_width = max(int(width_a), int(width_b))

    height_a = np.linalg.norm(tr - br)
    height_b = np.linalg.norm(tl - bl)
    max_height = max(int(height_a), int(height_b))

    if max_width < 50 or max_height < 50:
        return image

    dst = np.array(
        [
            [0, 0],
            [max_width - 1, 0],
            [max_width - 1, max_height - 1],
            [0, max_height - 1],
        ],
        dtype="float32",
    )

    matrix = cv2.getPerspectiveTransform(rect, dst)
    return cv2.warpPerspective(image, matrix, (max_width, max_height))


def smart_crop(image: np.ndarray) -> np.ndarray:
    original = image.copy()

    max_dim = 1400
    h, w = image.shape[:2]
    scale = min(1.0, max_dim / float(max(h, w)))
    resized = cv2.resize(image, None, fx=scale, fy=scale) if scale < 1 else image.copy()

    gray = cv2.cvtColor(resized, cv2.COLOR_BGR2GRAY)
    gray = cv2.GaussianBlur(gray, (5, 5), 0)
    edges = cv2.Canny(gray, 50, 150)
    edges = cv2.dilate(edges, None, iterations=1)

    contours, _ = cv2.findContours(edges, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
    contours = sorted(contours, key=cv2.contourArea, reverse=True)[:10]

    page = None
    for contour in contours:
        perimeter = cv2.arcLength(contour, True)
        approx = cv2.approxPolyDP(contour, 0.02 * perimeter, True)
        if len(approx) == 4:
            area = cv2.contourArea(approx)
            if area > resized.shape[0] * resized.shape[1] * 0.20:
                page = approx.reshape(4, 2).astype("float32")
                break

    if page is None:
        return original

    page /= scale
    return four_point_transform(original, page)


def rotate_by_osd(image: np.ndarray) -> np.ndarray:
    try:
        rgb = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
        osd = pytesseract.image_to_osd(rgb, output_type=pytesseract.Output.DICT)
        rotate = int(osd.get("rotate", 0))
        if rotate == 90:
            return cv2.rotate(image, cv2.ROTATE_90_CLOCKWISE)
        if rotate == 180:
            return cv2.rotate(image, cv2.ROTATE_180)
        if rotate == 270:
            return cv2.rotate(image, cv2.ROTATE_90_COUNTERCLOCKWISE)
    except Exception:
        pass
    return image


def enhance_document(image: np.ndarray) -> np.ndarray:
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)

    # Gentle denoise keeps text edges while reducing camera noise.
    gray = cv2.fastNlMeansDenoising(gray, None, 8, 7, 21)

    # CLAHE improves uneven lighting without destroying all grayscale detail.
    clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
    enhanced = clahe.apply(gray)

    # Convert back to BGR so JPEG and OCR handling remain simple.
    return cv2.cvtColor(enhanced, cv2.COLOR_GRAY2BGR)


def run_ocr(image: np.ndarray, language: str) -> str:
    rgb = cv2.cvtColor(image, cv2.COLOR_BGR2RGB)
    config = "--oem 3 --psm 6"
    return pytesseract.image_to_string(rgb, lang=language, config=config).strip()


def decode_upload(raw: bytes) -> np.ndarray:
    arr = np.frombuffer(raw, np.uint8)
    image = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if image is None:
        raise HTTPException(status_code=400, detail="Unsupported or invalid image.")
    return image


@app.get("/health")
def health():
    return {"status": "ok", "service": "grd_job_scan"}


def process_image(raw: bytes, language: str):
    """Crop, orient, enhance and OCR an image. Returns (enhanced_image, text)."""
    image = decode_upload(raw)

    cropped = smart_crop(image)
    oriented = rotate_by_osd(cropped)
    enhanced = enhance_document(oriented)

    try:
        text = run_ocr(enhanced, language)
    except pytesseract.TesseractNotFoundError as exc:
        raise HTTPException(
            status_code=500,
            detail="Tesseract OCR is not installed on the backend server.",
        ) from exc
    except pytesseract.TesseractError as exc:
        raise HTTPException(
            status_code=500,
            detail=f"OCR failed. Check installed language packs. Details: {exc}",
        ) from exc
    return enhanced, text


async def read_limited(file: UploadFile) -> bytes:
    raw = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File is larger than 20 MB.")
    return raw


@app.post("/api/scan")
async def scan_document(
    file: UploadFile = File(...),
    language: str = Form("eng"),
    user: str = Depends(current_user),
):
    if not (file.content_type or "").startswith("image/"):
        raise HTTPException(status_code=400, detail="Please upload an image file.")

    enhanced, text = process_image(await read_limited(file), language)

    scan_id = str(uuid.uuid4())
    output_path = SCAN_DIR / f"{scan_id}.jpg"
    cv2.imwrite(str(output_path), enhanced, [int(cv2.IMWRITE_JPEG_QUALITY), 92])

    return {
        "scan_id": scan_id,
        "text": text,
        "image_url": f"/api/scans/{scan_id}/image",
    }


@app.post("/api/parse")
async def parse_file(
    file: UploadFile = File(...),
    language: str = Form("eng"),
    user: str = Depends(current_user),
):
    """Extract text from an uploaded txt/csv, docx, pdf or image file."""
    name = file.filename or "upload"
    suffix = Path(name).suffix.lower()
    raw = await read_limited(file)

    if suffix in TEXT_EXTENSIONS:
        kind, text = "text", raw.decode("utf-8", errors="replace")
    elif suffix == ".docx":
        try:
            doc = Document(io.BytesIO(raw))
        except Exception as exc:
            raise HTTPException(status_code=400, detail="Could not read this .docx file.") from exc
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text.strip() for cell in row.cells))
        kind, text = "docx", "\n\n".join(parts)
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise HTTPException(
                status_code=501,
                detail="PDF parsing needs the 'pypdf' package on the server (pip install pypdf).",
            ) from exc
        try:
            reader = PdfReader(io.BytesIO(raw))
            text = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
        except Exception as exc:
            raise HTTPException(status_code=400, detail="Could not read this PDF file.") from exc
        kind = "pdf"
    elif suffix in IMAGE_EXTENSIONS or (file.content_type or "").startswith("image/"):
        _, text = process_image(raw, language)
        kind = "image"
    else:
        raise HTTPException(
            status_code=415,
            detail="Unsupported file type. Use txt, csv, docx, pdf or an image.",
        )

    text = text.strip()
    return {"filename": name, "type": kind, "text": text, "chars": len(text)}


@app.get("/api/scans/{scan_id}/image")
def get_scan_image(scan_id: str):
    path = SCAN_DIR / f"{scan_id}.jpg"
    if not path.exists():
        raise HTTPException(status_code=404, detail="Scan not found.")
    return FileResponse(path, media_type="image/jpeg", filename=f"{scan_id}.jpg")


@app.delete("/api/scans/{scan_id}")
def delete_scan(scan_id: str, user: str = Depends(current_user)):
    path = SCAN_DIR / f"{scan_id}.jpg"
    if path.exists():
        path.unlink()
    return {"deleted": True, "scan_id": scan_id}


@app.post("/api/export-word")
def export_word(request: ExportRequest, user: str = Depends(current_user)):
    if not request.pages:
        raise HTTPException(status_code=400, detail="At least one page is required.")

    document = Document()
    document.core_properties.title = request.title

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    document.add_heading(request.title, level=1)

    for index, page in enumerate(request.pages, start=1):
        if len(request.pages) > 1:
            document.add_heading(f"Page {index}", level=2)

        # Keep output editable: OCR text is inserted as real Word paragraphs.
        chunks = [p.strip() for p in page.text.replace("\r\n", "\n").split("\n\n") if p.strip()]
        if not chunks and page.text.strip():
            chunks = [page.text.strip()]

        for chunk in chunks:
            document.add_paragraph(chunk)

        if index < len(request.pages):
            document.add_page_break()

    export_id = str(uuid.uuid4())
    safe_name = "".join(c if c.isalnum() or c in (" ", "-", "_") else "_" for c in request.title).strip()
    safe_name = safe_name or "GRD_Job_Scan"
    filename = f"{safe_name}_{export_id[:8]}.docx"
    path = EXPORT_DIR / filename
    document.save(path)

    return {
        "export_id": export_id,
        "filename": filename,
        "download_url": f"/api/exports/{filename}",
    }


@app.get("/api/exports/{filename}")
def download_export(filename: str):
    safe_path = EXPORT_DIR / Path(filename).name
    if not safe_path.exists():
        raise HTTPException(status_code=404, detail="Export not found.")
    return FileResponse(
        safe_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=safe_path.name,
    )
